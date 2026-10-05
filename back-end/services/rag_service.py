# ── LlamaIndex core ───────────────────────────────────────────────────────────
import io
import logging

from config import (
    OLLAMA_BASE_URL,
    OLLAMA_EMBEDDING_MODEL,
    OLLAMA_MODEL,
    QDRANT_COLLECTION_NAME,
    QDRANT_URL,
)
from llama_index.core import (
    Document,
    PromptTemplate,
    Settings,
    StorageContext,
    VectorStoreIndex,
)
from llama_index.core.chat_engine import CondensePlusContextChatEngine
from llama_index.core.llms import ChatMessage, MessageRole
from llama_index.core.memory import ChatMemoryBuffer
from llama_index.core.node_parser import SemanticSplitterNodeParser
from llama_index.embeddings.ollama import OllamaEmbedding
from llama_index.llms.ollama import Ollama

# ── LlamaIndex integrations ───────────────────────────────────────────────────
from llama_index.vector_stores.qdrant import QdrantVectorStore

# ── Qdrant ────────────────────────────────────────────────────────────────────
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams

logger = logging.getLogger(__name__)


# ============================================================================
# STEP 1: Connect LlamaIndex to Ollama
# ============================================================================
logger.info("Connecting to Ollama for LLM and embeddings...")


# Configure embedding model (converts text to vectors)
Settings.embed_model = OllamaEmbedding(
    model_name=OLLAMA_EMBEDDING_MODEL,
    base_url=OLLAMA_BASE_URL,
)
logger.info(f"Embeddings: {OLLAMA_EMBEDDING_MODEL} at {OLLAMA_BASE_URL}")


# Configure LLM (generates answers)
Settings.llm = Ollama(
    model=OLLAMA_MODEL,
    base_url=OLLAMA_BASE_URL,
    request_timeout=90.0,
)
logger.info(f"LLM: {OLLAMA_MODEL} at {OLLAMA_BASE_URL}")


# ============================================================================
# STEP 2: Connect to Qdrant Vector Database
# ============================================================================
logger.info("Connecting to Qdrant vector database...")

qdrant_client = QdrantClient(url=QDRANT_URL)
logger.info(f"Qdrant: {QDRANT_URL}")

COLLECTION_NAME = QDRANT_COLLECTION_NAME


def _collection_exists() -> bool:
    """Return True if the configured Qdrant collection exists."""
    try:
        collections = qdrant_client.get_collections().collections
        return any(c.name == COLLECTION_NAME for c in collections)
    except Exception as e:
        logger.error(f"Failed to check Qdrant collections: {e}")
        return False


# ============================================================================
# STEP 3: Helper to Initialize Splitters with Ollama
# ============================================================================
def create_semantic_splitter(buffer_size=1, breakpoint_percentile_threshold=95):
    """
    Factory function to create SemanticSplitterNodeParser with Ollama embeddings.
    Bypasses OpenAI default by explicitly passing embed_model.
    """
    return SemanticSplitterNodeParser(
        embed_model=Settings.embed_model,
        buffer_size=buffer_size,
        breakpoint_percentile_threshold=breakpoint_percentile_threshold,
    )


# ============================================================================
# STEP 4: Configure Intelligent Text Chunking
# ============================================================================
text_splitter = create_semantic_splitter()
logger.info("Text splitter configured with semantic chunking (Ollama embeddings)")


# ============================================================================
# STEP 5: Initialize LlamaIndex with Qdrant Vector Store
# ============================================================================
def initialize_index():
    """
    Load existing vector store or return None if collection does not exist yet.
    Collection is auto-created when first documents are added.
    """
    if not _collection_exists():
        logger.info(
            f"Collection '{COLLECTION_NAME}' will be created on first document upload"
        )
        return None

    try:
        # Connect LlamaIndex to existing Qdrant collection
        vector_store = QdrantVectorStore(
            client=qdrant_client, collection_name=COLLECTION_NAME
        )
        index = VectorStoreIndex.from_vector_store(vector_store)
        logger.info(f"Loaded index from collection '{COLLECTION_NAME}'")
        return index
    except Exception as e:
        logger.error(f"Error initializing index: {e}")
        return None


# Global index instance
index = initialize_index()


# ============================================================================
# File text extraction
# ============================================================================

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


def get_file_extension(filename: str) -> str:
    """Return the lowercased extension (with leading dot) of a filename, or ''."""
    return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def extract_text_from_file(filename: str, content: bytes) -> str:
    """Extract plain text from a file given its name and raw bytes."""

    ext = get_file_extension(filename)

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext}'. "
            f"Allowed: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    # ── Plain text / Markdown ─────────────────────────────────────────────
    if ext in (".txt", ".md"):
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        text = text.strip()
        if not text:
            raise ValueError("The text file appears to be empty.")
        return text

    # ── PDF ───────────────────────────────────────────────────────────────
    if ext == ".pdf":
        try:
            import pypdf
        except ImportError:
            raise RuntimeError(
                "pypdf is not installed. Run: pip install pypdf"
            ) from None

        reader = pypdf.PdfReader(io.BytesIO(content))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        pages = [p for p in pages if p]

        if not pages:
            raise ValueError(
                "Could not extract text from PDF. "
                "It may be a scanned image — OCR is not supported yet."
            )
        return "\n\n".join(pages)

    # ── DOCX ──────────────────────────────────────────────────────────────
    if ext == ".docx":
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise RuntimeError(
                "python-docx is not installed. Run: pip install python-docx"
            ) from None

        doc = DocxDocument(io.BytesIO(content))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        if not paragraphs:
            raise ValueError("Could not extract text from DOCX file.")
        return "\n\n".join(paragraphs)

    raise ValueError(f"Unhandled extension: {ext}")


# ============================================================================
# Core RAG Functions
# ============================================================================


def add_documents(texts: list[str]) -> dict:
    """Ingest a list of raw text strings into the vector store."""

    global index
    try:
        logger.info(f"Processing {len(texts)} document(s)...")

        # Create documents from raw text
        documents = [Document(text=text) for text in texts]

        # Create the collection explicitly if it does not exist yet
        if not _collection_exists():
            logger.info("Creating new Qdrant collection...")
            qdrant_client.create_collection(
                collection_name=COLLECTION_NAME,
                vectors_config=VectorParams(size=768, distance=Distance.COSINE),
            )
            logger.info(f"Collection '{COLLECTION_NAME}' created")

        # Create vector store
        vector_store = QdrantVectorStore(
            client=qdrant_client, collection_name=COLLECTION_NAME
        )
        storage_context = StorageContext.from_defaults(vector_store=vector_store)

        # Chunk once, then append to the existing index or build a new one
        nodes = text_splitter.get_nodes_from_documents(documents)
        logger.info(f"  -> Split into {len(nodes)} chunks")

        if index is not None:
            logger.info("Adding to existing collection...")
            index.insert_nodes(nodes)
        else:
            logger.info("Creating index...")
            index = VectorStoreIndex(nodes, storage_context=storage_context)

        chunks_created = len(nodes)
        total_chunks = qdrant_client.get_collection(COLLECTION_NAME).points_count
        logger.info(
            f"Created {chunks_created} chunk(s); collection now has {total_chunks}"
        )

        return {
            "status": "success",
            "documents_added": len(documents),
            "chunks_created": chunks_created,
            "total_chunks": total_chunks,
            "message": f"Successfully split {len(documents)} document(s) into {chunks_created} optimized chunks",
        }
    except Exception as e:
        logger.error(f"Error adding documents: {e}")
        return {"status": "error", "message": str(e)}


def query_rag(question: str, similarity_top_k: int = 5) -> str:
    """Run a single-turn RAG query against the vector store."""

    global index
    try:
        if index is None:
            index = initialize_index()

        if index is None:
            return "No documents in the knowledge base. Please upload documents first."

        logger.info(f"Query: {question[:100]}...")
        query_engine = index.as_query_engine(similarity_top_k=similarity_top_k)
        response = query_engine.query(question)
        return str(response)
    except Exception as e:
        logger.error(f"Query error: {e}")
        return "Error: unable to process the query."


def get_collection_info() -> dict:
    """Get statistics about the Qdrant collection."""
    try:
        if not _collection_exists():
            return {
                "status": "success",
                "collection_name": COLLECTION_NAME,
                "points_count": 0,
                "message": "Collection does not exist yet",
            }

        info = qdrant_client.get_collection(COLLECTION_NAME)
        return {
            "status": "success",
            "collection_name": COLLECTION_NAME,
            "points_count": info.points_count,
            "vectors_in_collection": info.points_count,
        }
    except Exception as e:
        return {"status": "error", "message": str(e)}


def reset_collection() -> dict:
    """Delete all documents from the collection."""
    global index
    try:
        qdrant_client.delete_collection(COLLECTION_NAME)
        index = None
        logger.info("Collection reset")
        return {"status": "success", "message": "Collection deleted successfully"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


# ============================================================================
# Prompt templates (French — intentional, model responds to users in French)
# ============================================================================

CONDENSE_PROMPT = PromptTemplate(
    "Étant donné l'historique de conversation ci-dessous et une question de suivi, "
    "reformulez la question de suivi en une question autonome qui capture tout le contexte nécessaire. "
    "Répondez uniquement avec la question reformulée, sans explication.\n\n"
    "Historique de conversation :\n"
    "{chat_history}\n\n"
    "Question de suivi : {question}\n"
    "Question autonome :"
)

_CONTEXT_PROMPT_BODY = (
    "Tu aides les employés dans leurs tâches professionnelles quotidiennes.\n\n"
    "Règles de comportement :\n"
    "- Réponds toujours en français, avec un ton professionnel mais accessible.\n"
    "- Tu as accès à l'historique complet de la conversation : utilise les informations précédentes fournies par l'utilisateur pour maintenir la cohérence et la continuité. Garde en mémoire les contextes, les détails spécifiques et les préférences mentionnés antérieurement.\n"
    "- Si l'utilisateur te demande quelles informations tu as ou quelles sont tes sources disponibles, réponds de manière vague et demande des précisions sur le sujet spécifique qui l'intéresse. Ne divulgue PAS la liste complète des documents ou sources à ta disposition. Exemple : 'Je dispose de plusieurs documents internes. Quel sujet spécifique t'intéresse ?'\n"
    "- Si le contexte fourni contient la réponse, utilise-le en priorité et cite les informations pertinentes sans mentionner les noms des fichiers sources.\n"
    "- Si le contexte est partiellement utile, combine-le avec tes connaissances générales et l'historique de conversation.\n"
    "- Si la question est hors sujet professionnel (divertissement, opinion personnelle, etc.), "
    "réponds brièvement et invite poliment l'utilisateur à revenir sur des sujets liés à son travail ou "
    "aux documents internes de la base de connaissances.\n"
    "- Ne fabrique jamais d'informations techniques ou réglementaires : "
    "si tu n'es pas sûr, dis-le clairement et suggère de consulter la documentation officielle.\n\n"
    "Contexte de la base de connaissances :\n"
    "---------------------\n"
    "{context_str}\n"
    "---------------------\n\n"
    "Question : {query_str}\n"
    "Réponse :"
)


def _context_prompt(persona_prompt: str | None = None) -> PromptTemplate:
    """Build the context prompt, optionally with a persona instruction."""
    persona = (
        f"Adopte la personnalité suivante : {persona_prompt}\n"
        if persona_prompt
        else ""
    )
    return PromptTemplate(
        "Tu es RAG PoC L3, l'assistant intelligent. " + persona + _CONTEXT_PROMPT_BODY
    )


# ============================================================================
# History conversion
# ============================================================================


def _to_chat_messages(history: list) -> list[ChatMessage]:
    """Convert SQLite history [{role, content}] to LlamaIndex ChatMessage list."""
    role_map = {
        "user": MessageRole.USER,
        "assistant": MessageRole.ASSISTANT,
    }
    return [
        ChatMessage(
            role=role_map.get(msg.get("role", "user"), MessageRole.USER),
            content=msg["content"],
        )
        for msg in history
    ]


# ============================================================================
# Core chat function
# ============================================================================
def generate_session_title(question: str, answer: str) -> str:
    """Generate a short session title from the first question/answer pair."""
    try:
        prompt_msg = ChatMessage(
            role=MessageRole.USER,
            content=(
                "Génère un titre très court (3 à 6 mots maximum) qui résume "
                "l'idée générale de cet échange. Le titre doit être nominal "
                "(pas de verbe conjugué), dans la même langue que la question. "
                "Retourne UNIQUEMENT le titre, sans guillemets, sans ponctuation finale.\n\n"
                f"Question : {question}\n"
                f"Réponse : {answer[:300]}"
            ),
        )
        response = Settings.llm.chat([prompt_msg])
        title = str(response.message.content).strip().strip('"').strip("'")
        words = title.split()
        if len(words) > 8:
            title = " ".join(words[:6]) + "…"
        return title if title else " ".join(question.split()[:6])
    except Exception as e:
        logger.error(f"Title generation error: {e}")
        words = question.split()
        return " ".join(words[:6]) + ("…" if len(words) > 6 else "")


def generate_chat_response(
    question: str,
    history: list,
    similarity_top_k: int = 5,
    persona_prompt: str | None = None,
) -> str:
    """Generate a chat response using RAG if a knowledge base is available, otherwise fall back to plain LLM."""

    global index

    try:
        chat_history = _to_chat_messages(history)

        # Case 1: knowledge base available ─────────────────────────────
        if _collection_exists() and (
            index is not None or (index := initialize_index()) is not None
        ):
            retriever = index.as_retriever(similarity_top_k=similarity_top_k)
            retrieved = retriever.retrieve(question)
            logger.info(f"Retrieved {len(retrieved)} chunks for context")

            memory = ChatMemoryBuffer.from_defaults(
                chat_history=chat_history,
                token_limit=8000,
            )

            chat_engine = CondensePlusContextChatEngine.from_defaults(
                retriever=retriever,
                memory=memory,
                llm=Settings.llm,
                condense_question_prompt=CONDENSE_PROMPT,
                context_prompt=_context_prompt(persona_prompt),
                verbose=False,
            )

            response = chat_engine.chat(question)
            return str(response).strip()

        # Case 2: no knowledge base yet — plain LLM ────────────────────
        logger.info("No index available — falling back to plain LLM chat.")

        system_content = (
            "Tu es RAG PoC L3, l'assistant intelligent. "
            "Réponds en français, avec un ton professionnel et accessible. "
            "Aucune base de connaissances n'est encore disponible ; "
            "réponds à partir de tes connaissances générales et invite l'utilisateur "
            "à charger des documents pour obtenir des réponses plus précises."
        )
        if persona_prompt:
            system_content += f" Adopte la personnalité suivante : {persona_prompt}"

        system_msg = ChatMessage(
            role=MessageRole.SYSTEM,
            content=system_content,
        )

        messages = (
            [system_msg]
            + chat_history
            + [ChatMessage(role=MessageRole.USER, content=question)]
        )
        llm_response = Settings.llm.chat(messages)
        return str(llm_response.message.content).strip()

    except Exception as e:
        logger.error(f"Chat generation error: {e}")
        return "Error: unable to generate a response."
